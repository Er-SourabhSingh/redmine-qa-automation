"""
Crux-specific test-data registry generator — adapted from gen_testdata_xlsx.py's SHEETS
dict for redmineflux_crux's own entity types (users, projects, issues, agents, pipelines)
instead of Helpdesk's (organizations, customers, SLAs, ...). See CLAUDE.md section
"Test Data Registry (per environment)".

Usage:
    python scripts/gen_crux_testdata_xlsx.py <out_path.xlsx> <env_label>
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT_PATH = sys.argv[1]
ENV_LABEL = sys.argv[2]
PLUGIN_NAME = "Redmineflux Crux"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

STATUS_CHOICES = '"Active,Deactivated,Deleted,Unknown"'

SHEETS = {
    "README": {"columns": [], "rows": []},
    "Users": {
        "columns": ["Login", "Password", "Role(s)", "Project Membership", "Purpose", "Status",
                    "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [16, 12, 22, 18, 28, 14, 22, 14, 14, 40],
    },
    "Projects": {
        "columns": ["Identifier", "Name", "Public?", "Modules Enabled", "Trackers Assigned",
                    "Purpose", "Status", "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [16, 18, 10, 22, 18, 26, 14, 22, 14, 14, 40],
    },
    "Issues (fixtures)": {
        "columns": ["Issue #", "Subject / Marker", "Project", "Purpose / Used By (TC)",
                    "Status", "Created Date", "Last Verified", "Notes"],
        "widths": [10, 40, 16, 24, 14, 14, 14, 40],
    },
    "Agents (test)": {
        "columns": ["Agent ID", "Name", "Purpose / Used By (TC)", "Status",
                    "Created Date", "Last Verified", "Notes"],
        "widths": [22, 26, 24, 14, 14, 14, 40],
    },
    "Pipelines (test)": {
        "columns": ["Pipeline ID", "Name", "Purpose / Used By (TC)", "Status",
                    "Created Date", "Last Verified", "Notes"],
        "widths": [22, 26, 24, 14, 14, 14, 40],
    },
    "Seeded WPs Consumed": {
        "columns": ["WP ID", "Goal", "Gate Consumed", "Consumed By (TC / Session)",
                    "Date", "Notes"],
        "widths": [10, 40, 20, 22, 14, 40],
    },
}

wb = Workbook()
wb.remove(wb.active)

readme = wb.create_sheet("README")
readme.column_dimensions["A"].width = 100
readme_lines = [
    f"{PLUGIN_NAME} — Test Data Registry",
    f"Environment: {ENV_LABEL}",
    "",
    "PURPOSE",
    "Tracks every fixture (user, project, issue, agent, pipeline) created specifically for QA testing on",
    "this environment, plus any pre-existing SEEDED entity (e.g. a Work Package) whose state a test run",
    "permanently changed (e.g. approving one of its gates) — so later sessions know not to assume it is",
    "still in its original state.",
    "",
    "HOW TO USE",
    "- Per QA_CREDENTIALS.md Authentication Rule #1: always use the shared seed-user pool there",
    "  (luna.blossom, daisy.skye, ...) for any new test account on ANY environment, including this",
    "  plugin's own dedicated instance — do not invent ad hoc usernames.",
    "- Before creating a new fixture, check the relevant sheet first to see if one already exists that",
    "  can be reused instead of creating a duplicate.",
    "- After a session creates, modifies, or deletes a fixture, update its row here immediately.",
    "- The 'Seeded WPs Consumed' sheet is specifically for pre-existing crux-core seed data",
    "  (not created by QA) whose state a test permanently changed, e.g. a gate approval that cannot be",
    "  reversed via the UI — so later suites don't assume that WP is still in a pristine/pending state.",
    "",
    "STATUS COLUMN LEGEND",
    "- Active      : fixture exists and is in normal use",
    "- Deactivated : exists but disabled/paused/retired",
    "- Deleted     : removed from the environment (kept as a row here for history)",
    "- Unknown     : not verified recently — re-check before relying on it",
    "",
    "ONE WORKBOOK PER ENVIRONMENT",
    "This workbook is for the local Crux-dedicated Docker instance (localhost:3014) specifically — a",
    "different, plugin-dedicated instance from QA_CREDENTIALS.md's default localhost:3006, per that",
    "file's own note that plugins with a dedicated instance document their users/data in their own docs.",
]
for i, line in enumerate(readme_lines, start=1):
    cell = readme.cell(row=i, column=1, value=line)
    if i == 1:
        cell.font = Font(bold=True, size=14)
    elif line.isupper() and line != "":
        cell.font = Font(bold=True, size=11)

for name, spec in SHEETS.items():
    if name == "README":
        continue
    ws = wb.create_sheet(name)
    cols = spec["columns"]
    widths = spec["widths"]
    for c_idx, header in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(c_idx)].width = widths[c_idx - 1]
    ws.freeze_panes = "A2"

    status_col_name = "Status" if "Status" in cols else None
    if status_col_name:
        status_col_idx = cols.index(status_col_name) + 1
        col_letter = get_column_letter(status_col_idx)
        dv = DataValidation(type="list", formula1=STATUS_CHOICES, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}500")

wb.save(OUT_PATH)
print("Wrote", OUT_PATH)
