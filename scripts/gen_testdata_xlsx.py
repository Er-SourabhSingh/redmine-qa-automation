"""
Generate a per-environment test-data registry workbook (real .xlsx) for a plugin's
automation/testdata/ folder — see CLAUDE.md section "Test Data Registry (per environment)".

Usage:
    pip install openpyxl   (one-time)
    python scripts/gen_testdata_xlsx.py <out_path.xlsx> <env_label> [plugin_display_name]

Example:
    python scripts/gen_testdata_xlsx.py \
        plugins/redmineflux_helpdesk_qa/automation/testdata/HELPDESK_TESTDATA_FORGE.xlsx \
        "Forge (https://flux-xxxxx.forge.zehntech.com/, per QA_CREDENTIALS_FORGE.md)" \
        "Redmineflux Helpdesk"

Re-run with the same out_path to reset a workbook to a blank template (this always regenerates
from scratch — it does NOT merge with an existing file's data rows). Edit the SHEETS dict below
to change which entity sheets/columns a given plugin's registry needs.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT_PATH = sys.argv[1]
ENV_LABEL = sys.argv[2]  # e.g. "Local (http://localhost:3006)"
PLUGIN_NAME = sys.argv[3] if len(sys.argv) > 3 else "Plugin"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")

STATUS_CHOICES = '"Active,Deactivated,Deleted,Unknown"'

SHEETS = {
    "README": {
        "columns": [],
        "rows": [],
    },
    "Organizations": {
        "columns": ["Name", "Status", "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [24, 14, 22, 14, 14, 40],
    },
    "Customers": {
        "columns": ["Login", "Email", "Organization", "Project", "SLA", "Support Level", "Status",
                    "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [18, 22, 18, 18, 16, 14, 14, 22, 14, 14, 40],
    },
    "SLAs": {
        "columns": ["Name", "Project (blank = global)", "Response Time", "Resolution Time",
                    "Working Hours", "Working Days", "Active", "Status",
                    "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [18, 22, 14, 16, 16, 16, 8, 14, 22, 14, 14, 40],
    },
    "Support Levels": {
        "columns": ["Name", "Project", "Level Order", "Assignees", "Escalates To", "Status",
                    "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [16, 18, 12, 26, 16, 14, 22, 14, 14, 40],
    },
    "Products": {
        "columns": ["Name", "Code", "Project", "Status", "Created By (TC / Session)",
                    "Created Date", "Last Verified", "Notes"],
        "widths": [18, 14, 18, 14, 22, 14, 14, 40],
    },
    "Support Packages": {
        "columns": ["Name", "Status", "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [22, 14, 22, 14, 14, 40],
    },
    "Canned Responses": {
        "columns": ["Name", "Status", "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [22, 14, 22, 14, 14, 40],
    },
    "Holidays": {
        "columns": ["Name", "Start Date", "End Date", "Status", "Created By (TC / Session)",
                    "Created Date", "Last Verified", "Notes"],
        "widths": [24, 14, 14, 14, 22, 14, 14, 40],
    },
    "Tickets (fixtures)": {
        "columns": ["Ticket #", "Subject", "Project", "Current Status", "Purpose / Used By (TC)",
                    "Created Date", "Last Verified", "Notes"],
        "widths": [10, 30, 18, 16, 24, 14, 14, 40],
    },
    "Prepaid Budgets": {
        "columns": ["Organization", "Project", "Approved (h)", "Used (h)", "Remaining (h)",
                    "Run-out Mode", "Created By (TC / Session)", "Created Date", "Last Verified", "Notes"],
        "widths": [18, 18, 12, 10, 12, 14, 22, 14, 14, 40],
    },
}

wb = Workbook()
wb.remove(wb.active)

readme = wb.create_sheet("README")
readme.column_dimensions["A"].width = 100
readme_lines = [
    f"{PLUGIN_NAME} — Test Data Registry",
    f"Environment: {ENV_LABEL}",
    "",
    "PURPOSE",
    "This workbook tracks what test data currently exists on THIS specific environment, so that when a",
    "test case's expected result depends on a name/ID that isn't fixed (customer, organization, SLA,",
    "support level, product, canned response, holiday, ticket, prepaid budget), we have a real baseline",
    "to check against instead of assuming or hardcoding a value that only held on a different run/server.",
    "",
    "HOW TO USE",
    "- Before creating a new fixture (e.g. a customer or SLA) for a test case, check the relevant sheet",
    "  first to see if one already exists with the name/config you need — reuse it instead of duplicating.",
    "  Several helpdesk entities refuse duplicate names (organizations, SLAs, support levels, products,",
    "  canned responses, holidays) — this registry is also how we avoid hitting that refusal.",
    "- After a test run that creates, modifies, or deletes an entity, update its row here immediately —",
    "  do not batch this for later. A stale registry is worse than no registry.",
    "- When a test case's expected result depends on the identity of an existing record (e.g. \"the SLA",
    "  named Standard\"), verify against this file first — it is the source of truth for this environment,",
    "  not the test case file (which may have been written against a different run or environment).",
    "",
    "STATUS COLUMN LEGEND",
    "- Active      : record exists and is in normal use",
    "- Deactivated : record exists but its Active flag is off (per the plugin's own deactivate feature)",
    "- Deleted     : record has been removed from the environment (kept as a row here for history)",
    "- Unknown     : not verified recently — re-check before relying on it",
    "",
    "NOTES ON TICKET IDs SPECIFICALLY",
    "Ticket numbers auto-increment and are NOT stable across environments or even across reseeds of the",
    "same environment. Do not hardcode a ticket # in a test case's expected result. Instead, record fixture",
    "tickets used repeatedly across sessions in the 'Tickets (fixtures)' sheet by their current # AND a",
    "stable description (subject/purpose), and re-verify the # is still correct before each session that",
    "depends on it.",
    "",
    "ONE WORKBOOK PER ENVIRONMENT",
    "A separate copy of this file exists per server (e.g. HELPDESK_TESTDATA_LOCAL.xlsx,",
    "HELPDESK_TESTDATA_FORGE.xlsx) — data on one environment says nothing about another.",
]
for i, line in enumerate(readme_lines, start=1):
    cell = readme.cell(row=i, column=1, value=line)
    if i == 1:
        cell.font = Font(bold=True, size=14)
    elif line.isupper() and line != "":
        cell.font = Font(bold=True, size=11)

for name, spec in SHEETS.items():
    if name == "README":
        continue
    ws = wb.create_sheet(name)
    cols = spec["columns"]
    widths = spec["widths"]
    for c_idx, header in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(c_idx)].width = widths[c_idx - 1]
    ws.freeze_panes = "A2"

    # Status data-validation dropdown, applied to a generous range for future rows
    status_col_name = "Status" if "Status" in cols else ("Active" if "Active" in cols else None)
    if status_col_name:
        status_col_idx = cols.index(status_col_name) + 1
        col_letter = get_column_letter(status_col_idx)
        dv = DataValidation(type="list", formula1=STATUS_CHOICES, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col_letter}2:{col_letter}500")

wb.save(OUT_PATH)
print("Wrote", OUT_PATH)
